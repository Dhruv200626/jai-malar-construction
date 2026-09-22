"""
Billing views — list, create, detail, edit, delete, Excel export
"""
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Q
from django.core.paginator import Paginator
from django.http import HttpResponse
from decimal import Decimal, InvalidOperation
import json

from .models import Bill, BillItem
from clients.models import Client
from suppliers.models import Supplier


# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────

def _parse_decimal(value, default=Decimal('0')):
    try:
        return Decimal(str(value)).quantize(Decimal('0.01'))
    except (InvalidOperation, TypeError, ValueError):
        return default


def _save_bill_items(bill, post):
    """Delete old items and save fresh ones from POST data."""
    bill.items.all().delete()
    descriptions = post.getlist('item_description')
    quantities   = post.getlist('item_quantity')
    rates        = post.getlist('item_rate')
    units        = post.getlist('item_unit')
    hsn_codes    = post.getlist('item_hsn')

    for i, desc in enumerate(descriptions):
        desc = desc.strip()
        if not desc:
            continue
        qty  = _parse_decimal(quantities[i] if i < len(quantities) else '0')
        rate = _parse_decimal(rates[i]       if i < len(rates)       else '0')
        if qty <= 0 or rate <= 0:
            continue
        BillItem.objects.create(
            bill=bill,
            description=desc,
            quantity=qty,
            rate=rate,
            unit=units[i].strip()    if i < len(units)     else '',
            hsn_code=hsn_codes[i].strip() if i < len(hsn_codes) else '',
        )


def _recalculate_and_save(bill):
    """Recalculate totals from items then save."""
    items = bill.items.all()
    subtotal = sum(i.amount for i in items)
    bill.subtotal        = subtotal
    discount_amount      = subtotal * (bill.discount_percent / 100)
    bill.discount_amount = round(discount_amount, 2)
    bill.taxable_amount  = round(subtotal - bill.discount_amount, 2)
    gst_total            = bill.taxable_amount * (bill.gst_percent / 100)
    bill.cgst_amount     = round(gst_total / 2, 2)
    bill.sgst_amount     = round(gst_total / 2, 2)
    bill.tax_amount      = round(gst_total, 2)
    bill.grand_total     = round(bill.taxable_amount + bill.tax_amount, 2)
    bill.save()


# ─────────────────────────────────────────────
# List / Search
# ─────────────────────────────────────────────

@login_required
def bill_list(request):
    qs = Bill.objects.select_related('client', 'supplier').order_by('-created_at')
    q          = request.GET.get('q', '').strip()
    client_id  = request.GET.get('client', '')
    status     = request.GET.get('status', '')

    if q:
        qs = qs.filter(
            Q(invoice_number__icontains=q) |
            Q(client__name__icontains=q)   |
            Q(client__company__icontains=q)|
            Q(supplier__name__icontains=q)
        )
    if client_id:
        qs = qs.filter(client_id=client_id)
    if status:
        qs = qs.filter(status=status)

    paginator = Paginator(qs, 20)
    page = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'billing/list.html', {
        'page_title': 'Bills & Invoices',
        'bills': page,
        'clients': Client.objects.filter(is_active=True).order_by('name'),
        'status_choices': Bill.STATUS_CHOICES,
        'q': q, 'client_id': client_id, 'status': status,
        'total_count': qs.count(),
    })


# ─────────────────────────────────────────────
# Create
# ─────────────────────────────────────────────

@login_required
def bill_create(request):
    clients   = Client.objects.filter(is_active=True).order_by('name')
    suppliers = Supplier.objects.filter(is_active=True).order_by('name')

    if request.method == 'POST':
        client_id   = request.POST.get('client')
        supplier_id = request.POST.get('supplier') or None
        bill_date   = request.POST.get('bill_date')
        due_date    = request.POST.get('due_date') or None
        status      = request.POST.get('status', 'draft')
        payment_method = request.POST.get('payment_method', 'cash')
        discount_percent = _parse_decimal(request.POST.get('discount_percent', '0'))
        gst_percent      = _parse_decimal(request.POST.get('gst_percent', '18'))
        notes  = request.POST.get('notes', '').strip()
        terms  = request.POST.get('terms', '').strip()

        if not client_id or not bill_date:
            messages.error(request, 'Client and bill date are required.')
            return render(request, 'billing/form.html', {
                'page_title': 'Create Bill', 'clients': clients,
                'suppliers': suppliers, 'post': request.POST,
            })

        if not request.POST.getlist('item_description') or not any(
            d.strip() for d in request.POST.getlist('item_description')
        ):
            messages.error(request, 'Add at least one item.')
            return render(request, 'billing/form.html', {
                'page_title': 'Create Bill', 'clients': clients,
                'suppliers': suppliers, 'post': request.POST,
            })

        bill = Bill(
            client_id=client_id,
            supplier_id=supplier_id,
            bill_date=bill_date,
            due_date=due_date,
            status=status,
            payment_method=payment_method,
            discount_percent=discount_percent,
            gst_percent=gst_percent,
            notes=notes,
            terms=terms,
            created_by=request.user,
        )
        bill.save()  # generates invoice number, sets financials to 0 first
        _save_bill_items(bill, request.POST)
        _recalculate_and_save(bill)

        messages.success(request, f'Bill {bill.invoice_number} created successfully.')
        return redirect('billing:detail', pk=bill.pk)

    return render(request, 'billing/form.html', {
        'page_title': 'Create Bill',
        'clients': clients,
        'suppliers': suppliers,
        'status_choices': Bill.STATUS_CHOICES,
        'payment_choices': Bill.PAYMENT_CHOICES,
    })


# ─────────────────────────────────────────────
# Detail / View
# ─────────────────────────────────────────────

@login_required
def bill_detail(request, pk):
    bill = get_object_or_404(Bill.objects.select_related('client', 'supplier', 'created_by'), pk=pk)
    items = bill.items.all()
    return render(request, 'billing/detail.html', {
        'page_title': f'Bill {bill.invoice_number}',
        'bill': bill,
        'items': items,
    })


# ─────────────────────────────────────────────
# Edit
# ─────────────────────────────────────────────

@login_required
def bill_edit(request, pk):
    bill      = get_object_or_404(Bill, pk=pk)
    clients   = Client.objects.filter(is_active=True).order_by('name')
    suppliers = Supplier.objects.filter(is_active=True).order_by('name')
    items     = list(bill.items.all())

    if request.method == 'POST':
        bill.client_id       = request.POST.get('client')
        bill.supplier_id     = request.POST.get('supplier') or None
        bill.bill_date       = request.POST.get('bill_date')
        bill.due_date        = request.POST.get('due_date') or None
        bill.status          = request.POST.get('status', 'draft')
        bill.payment_method  = request.POST.get('payment_method', 'cash')
        bill.discount_percent = _parse_decimal(request.POST.get('discount_percent', '0'))
        bill.gst_percent     = _parse_decimal(request.POST.get('gst_percent', '18'))
        bill.notes           = request.POST.get('notes', '').strip()
        bill.terms           = request.POST.get('terms', '').strip()
        bill.save()
        _save_bill_items(bill, request.POST)
        _recalculate_and_save(bill)
        messages.success(request, f'Bill {bill.invoice_number} updated.')
        return redirect('billing:detail', pk=bill.pk)

    return render(request, 'billing/form.html', {
        'page_title': f'Edit Bill {bill.invoice_number}',
        'bill': bill,
        'items': items,
        'clients': clients,
        'suppliers': suppliers,
        'status_choices': Bill.STATUS_CHOICES,
        'payment_choices': Bill.PAYMENT_CHOICES,
    })


# ─────────────────────────────────────────────
# Delete
# ─────────────────────────────────────────────

@login_required
def bill_delete(request, pk):
    bill = get_object_or_404(Bill, pk=pk)
    if request.method == 'POST':
        inv = bill.invoice_number
        bill.delete()
        messages.success(request, f'Bill {inv} deleted.')
        return redirect('billing:list')
    return render(request, 'billing/confirm_delete.html', {
        'page_title': 'Delete Bill',
        'bill': bill,
    })


# ─────────────────────────────────────────────
# Excel Export
# ─────────────────────────────────────────────

@login_required
def bill_export_excel(request, pk):
    bill  = get_object_or_404(Bill.objects.select_related('client', 'supplier'), pk=pk)
    items = bill.items.all()

    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, 'openpyxl is not installed.')
        return redirect('billing:detail', pk=pk)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Invoice'

    # ── styles ──
    hdr_font    = Font(bold=True, color='FFFFFF', size=11)
    hdr_fill    = PatternFill('solid', fgColor='1A6FC4')
    title_font  = Font(bold=True, size=16, color='1A6FC4')
    sub_font    = Font(bold=True, size=11)
    label_font  = Font(bold=True, size=10, color='374151')
    value_font  = Font(size=10)
    total_fill  = PatternFill('solid', fgColor='EFF6FF')
    grand_fill  = PatternFill('solid', fgColor='1A6FC4')
    grand_font  = Font(bold=True, color='FFFFFF', size=12)
    thin        = Side(style='thin', color='CBD5E1')
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    center      = Alignment(horizontal='center', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    wrap        = Alignment(wrap_text=True, vertical='top')

    def cell(row, col, value, font=None, fill=None, align=None, bord=None, num_fmt=None):
        c = ws.cell(row=row, column=col, value=value)
        if font:    c.font      = font
        if fill:    c.fill      = fill
        if align:   c.alignment = align
        if bord:    c.border    = bord
        if num_fmt: c.number_format = num_fmt
        return c

    # col widths
    col_widths = [6, 28, 10, 14, 14, 14, 14]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 1
    # ── Title ──
    ws.merge_cells(f'A{r}:G{r}')
    cell(r, 1, 'JAI MALHAR CONSTRUCTION', font=title_font, align=center)
    ws.row_dimensions[r].height = 30
    r += 1
    ws.merge_cells(f'A{r}:G{r}')
    cell(r, 1, 'TAX INVOICE', font=Font(bold=True, size=13), align=center)
    ws.row_dimensions[r].height = 20
    r += 2

    # ── Bill Meta (two columns) ──
    meta_left = [
        ('Invoice No',    bill.invoice_number),
        ('Bill Date',     str(bill.bill_date)),
        ('Due Date',      str(bill.due_date) if bill.due_date else '—'),
        ('Status',        bill.get_status_display()),
        ('Payment',       bill.get_payment_method_display()),
    ]
    meta_right = [
        ('Client',        str(bill.client)),
        ('Company',       bill.client.company or '—'),
        ('Phone',         bill.client.phone or '—'),
        ('GST (Client)',  bill.client.gst_number or '—'),
        ('Supplier',      str(bill.supplier) if bill.supplier else '—'),
    ]
    for i, ((lbl, val), (rlbl, rval)) in enumerate(zip(meta_left, meta_right)):
        cell(r + i, 1, lbl,  font=label_font)
        cell(r + i, 2, val,  font=value_font)
        cell(r + i, 4, rlbl, font=label_font)
        cell(r + i, 5, rval, font=value_font)
    r += len(meta_left) + 1

    # ── Items header ──
    headers = ['#', 'Description', 'Unit', 'Qty', 'Rate (₹)', 'Amount (₹)', 'HSN/SAC']
    for ci, h in enumerate(headers, 1):
        cell(r, ci, h, font=hdr_font, fill=hdr_fill, align=center, bord=border)
    ws.row_dimensions[r].height = 18
    r += 1

    # ── Items rows ──
    for idx, item in enumerate(items, 1):
        cell(r, 1, idx,               font=value_font, align=center,      bord=border)
        cell(r, 2, item.description,  font=value_font, align=wrap,         bord=border)
        cell(r, 3, item.unit or '—',  font=value_font, align=center,      bord=border)
        cell(r, 4, float(item.quantity), font=value_font, align=right_align, bord=border, num_fmt='#,##0.00')
        cell(r, 5, float(item.rate),  font=value_font, align=right_align, bord=border, num_fmt='#,##0.00')
        cell(r, 6, float(item.amount),font=value_font, align=right_align, bord=border, num_fmt='#,##0.00')
        cell(r, 7, item.hsn_code or '—', font=value_font, align=center,  bord=border)
        ws.row_dimensions[r].height = 16
        r += 1

    r += 1  # gap

    # ── Totals block ──
    totals = [
        ('Subtotal',          float(bill.subtotal)),
        (f'Discount ({bill.discount_percent}%)', -float(bill.discount_amount)),
        ('Taxable Amount',    float(bill.taxable_amount)),
        (f'CGST ({bill.gst_percent/2}%)', float(bill.cgst_amount)),
        (f'SGST ({bill.gst_percent/2}%)', float(bill.sgst_amount)),
    ]
    for lbl, val in totals:
        ws.merge_cells(f'D{r}:E{r}')
        cell(r, 4, lbl,  font=sub_font,   fill=total_fill, align=right_align, bord=border)
        cell(r, 6, val,  font=value_font, fill=total_fill, align=right_align, bord=border, num_fmt='#,##0.00')
        r += 1

    # Grand total
    ws.merge_cells(f'D{r}:E{r}')
    cell(r, 4, 'GRAND TOTAL', font=grand_font, fill=grand_fill, align=right_align, bord=border)
    cell(r, 6, float(bill.grand_total), font=grand_font, fill=grand_fill, align=right_align, bord=border, num_fmt='#,##0.00')
    ws.row_dimensions[r].height = 22
    r += 2

    # ── Notes / Terms ──
    if bill.notes:
        cell(r, 1, 'Notes:', font=label_font)
        r += 1
        ws.merge_cells(f'A{r}:G{r}')
        cell(r, 1, bill.notes, font=value_font, align=wrap)
        ws.row_dimensions[r].height = 30
        r += 2
    if bill.terms:
        cell(r, 1, 'Terms & Conditions:', font=label_font)
        r += 1
        ws.merge_cells(f'A{r}:G{r}')
        cell(r, 1, bill.terms, font=value_font, align=wrap)
        ws.row_dimensions[r].height = 30

    # ── Serve ──
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{bill.invoice_number}.xlsx"'
    wb.save(response)
    return response
